"""นำเข้าข้อมูลแม่พิมพ์จากไฟล์ Excel (.xlsm) -> web/public/data/*.json + data/app.db

- สต็อก: ชีท DIEP1, PUNCHP1, PUNCHP3, DIEP3 (REDRAWP2 เป็นชีทเปล่า ข้าม)
- ใบซ่อม: ชีท P4R1 (วันที่ในไฟล์เก็บเป็น ค.ศ.1969 เพราะพิมพ์ปี 2 หลัก "69"
  -> แปลงเป็น 2026 / พ.ศ.2569 โดยอัตโนมัติ)
- สำรอง app.db เป็น app.db.bak ก่อนเขียนทุกครั้ง, ลบแถว DEMO ทิ้ง

รันจากโฟลเดอร์ Store1:  python tools/import_excel.py
"""
import glob
import io
import json
import os
import re
import shutil
import sqlite3
import sys
from datetime import datetime, date

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PATH_DIR = os.path.join(ROOT, "path")
DATA_DIR = os.path.join(ROOT, "web", "public", "data")
DB_PATH = os.path.join(ROOT, "data", "app.db")

STOCK_SHEETS = ["DIEP1", "PUNCHP1", "PUNCHP3", "DIEP3"]

STATUS_MAP = {
    "สแปร์": "spare",
    "ใช้งานบนเครื่อง": "installed",
    "ส่งซ่อม": "repair",
    "ยกเลิกใช้งาน": "retired",
}
STATUS_TH = {"spare": "สแปร์", "installed": "ใช้งานบนเครื่อง", "repair": "ส่งซ่อม", "retired": "ยกเลิกใช้งาน"}

WORKNAME_MAP = {
    "DIECUTTERP1": "DIE CUTTER P1",
    "DIECUTTERP3": "DIE CUTTER P3",
    "PUNCH CUTTER P1": "PUNCH CUTTER P1",
    "PUNCH CUTTER P3": "PUNCH CUTTER P3",
}

# 1969 (ที่ Excel ตีความจาก "69") -> 2026 (พ.ศ.2569)
YEAR_FIX = 2026 - 1969


def fix_date(v):
    """คืน YYYY-MM-DD (แก้ปี 1969 -> 2026) หรือ None"""
    if v is None or v == "":
        return None
    if isinstance(v, (datetime, date)):
        y = v.year + (YEAR_FIX if v.year == 1969 else 0)
        return f"{y:04d}-{v.month:02d}-{v.day:02d}"
    s = str(v).strip()[:10]
    if len(s) == 10 and s[4] == "-" and s.startswith("1969"):
        return f"2026{s[4:]}"
    return s if len(s) == 10 else None


def find_excel():
    files = glob.glob(os.path.join(PATH_DIR, "*.xlsm"))
    if not files:
        raise SystemExit(f"ไม่พบไฟล์ .xlsm ใน {PATH_DIR}")
    return sorted(files)[0]


def parse_stock(wb):
    molds = []
    seen_codes = set()
    for sheet in STOCK_SHEETS:
        if sheet not in wb.sheetnames:
            print(f"  ! ข้าม {sheet} (ไม่มีชีท)")
            continue
        ws = wb[sheet]
        header = [str(ws.cell(row=1, column=c).value or "") for c in range(1, ws.max_column + 1)]

        def col(name_part):
            for i, h in enumerate(header):
                if name_part in h:
                    return i
            return None

        i_work = col("ชื่อชิ้นงาน")
        i_size = 2  # คอลัมน์ SIZE ตัวแรก (index 2)
        i_no = col("No.")
        i_status = col("สถานะ")
        i_machine = col("เครื่องที่ใช้")
        i_remark = col("หมายเหตุ")
        # มิติเสริม: เก็บตามชื่อหัวจริง (ยกเว้นคอลัมน์หลัก)
        # หมายเหตุ: มีหัว "SIZE" 2 คอลัมน์ — อันที่ 2 คือเกรด (พิเศษ/ธรรมดา) เก็บแยกเป็น grade
        dim_idx = {}
        seen_size = 0
        i_grade = None
        for i, h in enumerate(header):
            if h == "SIZE":
                seen_size += 1
                if seen_size == 2:
                    i_grade = i
                continue
            if h and h not in ("ลำดับ", "ชื่อชิ้นงาน", "No.", "สถานะ", "เครื่องที่ใช้", "หมายเหตุ"):
                dim_idx[h] = i

        count = 0
        block = 1   # เลขบล็อก SIZE ตามต้นฉบับ (ขึ้นบล็อกใหม่ทุกครั้งที่เจอแถวหัวตาราง)
        seq = 0     # ลำดับแถวในบล็อก (เหมือนคอลัมน์ "ลำดับ" ใน Excel)
        for r in range(1, ws.max_row + 1):
            v0 = ws.cell(row=r, column=1).value
            if isinstance(v0, str) and "ลำดับ" in v0:
                if count > 0:
                    block += 1  # หัวตารางซ้ำ = ขึ้นบล็อก SIZE ใหม่
                seq = 0
                continue
            if not isinstance(v0, int):
                continue  # แถวรวม/หัวข้อ/ว่าง
            work = str(ws.cell(row=r, column=(i_work + 1)).value or "").strip() if i_work is not None else ""
            size = str(ws.cell(row=r, column=3).value or "").strip()
            no = ws.cell(row=r, column=(i_no + 1)).value if i_no is not None else None
            status_th = str(ws.cell(row=r, column=(i_status + 1)).value or "").strip() if i_status is not None else ""
            machine = ws.cell(row=r, column=(i_machine + 1)).value if i_machine is not None else None
            remark = ws.cell(row=r, column=(i_remark + 1)).value if i_remark is not None else None
            if no is None or str(no).strip() == "":
                continue
            no_s = str(no).strip()
            status = STATUS_MAP.get(status_th, "spare")
            tooling = WORKNAME_MAP.get(work, work or sheet)
            code = f"{sheet}-{size}-{no_s}".replace(" ", "")
            if code in seen_codes:  # กันรหัสซ้ำ (No. ซ้ำข้าม SIZE)
                code = f"{code}-{r}"
            seen_codes.add(code)
            dims = {}
            for h, i in dim_idx.items():
                v = ws.cell(row=r, column=i + 1).value
                if v not in (None, ""):
                    dims[h] = v
            grade_v = ws.cell(row=r, column=i_grade + 1).value if i_grade is not None else None
            seq += 1
            molds.append({
                "code": code,
                "sheet": sheet,
                "block": block,          # บล็อก SIZE ที่ n ของชีท (ตามต้นฉบับ)
                "seq": seq,              # ลำดับในบล็อก (ตามต้นฉบับ)
                "tooling_type": tooling,
                "work_name": work,
                "size": size,
                "no": no_s,
                "grade": str(grade_v).strip() if grade_v not in (None, "") else None,
                "status": status,
                "status_th": STATUS_TH[status],
                "machine": str(machine).strip() if machine not in (None, "") else None,
                "remark": str(remark).strip() if remark not in (None, "") else None,
                "dims": dims,
            })
            count += 1
        print(f"  {sheet}: {count} ตัว")
    return molds


def parse_repairs(wb):
    if "P4R1" not in wb.sheetnames:
        print("  ! ไม่มีชีท P4R1", flush=True)
        return []
    ws = wb["P4R1"]
    # อ่านหัวเพื่อหาตำแหน่งคอลัมน์จริง (กันไฟล์ขยับคอลัมน์)
    header = [(str(ws.cell(row=1, column=c).value or "").strip()) for c in range(1, ws.max_column + 1)]

    def find(*names):
        for i, h in enumerate(header):
            if h in names:
                return i
        return None

    i_no = find("เลขที่ต้นสังกัด")
    i_job = find("รายการ")
    i_pat = find("รหัสแบบ")
    i_qty = find("จำนวน")
    i_unit = find("หน่วย")
    i_month = find("เดือน")
    i_prio = find("สถานะ")
    i_ordered = find("วันที่สั่ง")
    i_due = find("กำหนดเสร็จ")
    i_received = find("วันที่รับงาน")
    i_machine = find("เครื่องที่ใช้")
    i_requester = find("ชื่อผู้สั่ง")
    i_receiver = find("ผู้รับงาน")
    i_detail = find("รายละเอียด")
    i_size = find("SIZE")
    repairs = []
    for r in range(2, ws.max_row + 1):
        no_v = ws.cell(row=r, column=(i_no + 1)).value if i_no is not None else None
        if not no_v or str(no_v).strip() in ("", "เลขที่ต้นสังกัด"):
            continue
        source_no = str(no_v).strip()
        if not source_no.startswith("P4R"):
            continue
        gv = lambda i: ws.cell(row=r, column=i + 1).value if i is not None else None
        job = str(gv(i_job) or "").strip()
        priority = str(gv(i_prio) or "").strip()
        ordered = fix_date(gv(i_ordered))
        due = fix_date(gv(i_due))
        received = fix_date(gv(i_received))
        # มีวันที่รับงานแล้ว = ดำเนินการเสร็จ/รับแล้ว, ไม่งั้น = ค้าง
        status = "received" if received else "open"
        month_v = gv(i_month)
        repairs.append({
            "source_no": source_no,
            "job_name": job,
            "pattern_code": str(gv(i_pat)).strip() if gv(i_pat) else None,
            "qty": str(gv(i_qty)).strip() if gv(i_qty) not in (None, "") else None,
            "unit": str(gv(i_unit)).strip() if gv(i_unit) not in (None, "") else None,
            "priority": priority or None,
            "status": status,
            "ordered_date": ordered,
            "due_date": due,
            "received_date": received,
            "month": int(str(month_v).strip()) if month_v is not None and str(month_v).strip().isdigit() else None,
            "machine": str(gv(i_machine)).strip() if gv(i_machine) else None,
            "requester": str(gv(i_requester)).strip() if gv(i_requester) else None,
            "receiver": str(gv(i_receiver)).strip() if gv(i_receiver) else None,
            "detail": str(gv(i_detail)).strip() if gv(i_detail) else None,
            "size_ref": str(gv(i_size)).strip() if gv(i_size) else None,
        })
    print(f"  P4R1: {len(repairs)} ใบ (ค้าง {sum(1 for r in repairs if r['status']=='open')} / รับแล้ว {sum(1 for r in repairs if r['status']=='received')})")
    return repairs


# ชื่องาน (ลบช่องว่าง/ตัวพิมพ์) -> ชีทสต็อก ; งานที่ไม่มีใน 4 ชีท = นอกทะเบียน
# หมายเหตุ: ใช้ exact match กับแต่ละส่วนที่แยกด้วย "+" (งานคู่) — ห้าม contains-match
# เพราะชื่อคล้ายกันอาจเป็นคนละชิ้นงาน (เช่น ดุม PUNCH CUTTER P3 ไม่ใช่ punch cutter)
JOB_TO_SHEET = {
    "DIECUTTERP1": "DIEP1",
    "PUNCHCUTTERP1": "PUNCHP1",
    "PUNCHCUTTERP3": "PUNCHP3",
    "DIECUTTERP3": "DIEP3",
    "PUNCHCUTERP1": "PUNCHP1",      # ตัวสะกดเพี้ยนในไฟล์
    "แก้PUNCHCUTTERP3": "PUNCHP3",   # prefix แก้ = งานซ่อมชนิดนั้น
}

# ชื่องาน (normalized) -> หมวดหมู่ (เช็คก่อน keyword)
JOB_GROUP_EXPLICIT = {
    "": "ไม่ระบุ",
    "สลัก": "สลัก",
    "RULLEY4นิ้ว": "พูลเลย์",
    "ฐานเครื่องมือวัดเคาท์เตอร์ซิงค์": "เครื่องมือวัด",
}

# keyword (ตามลำดับ) -> หมวดหมู่ ; ใช้กับ normalized key
GROUP_KEYWORDS = [
    ("PRESSURERING", "PRESSURE RING"),
    ("REDRAW", "REDRAW"),
    ("แหวนรอง", "แหวนรอง"),
    ("CENTERCORE", "CORE"),
    ("DIECORE", "CORE"),
    ("PUNCHHOLDER", "HOLDER"),
    ("HOLDER", "HOLDER"),
    ("TRIMMING", "TRIMMING"),
    ("BLANK", "BLANK"),
    ("BOTTOMKNIFE", "KNIFE"),
    ("KNIFE", "KNIFE"),
    ("LOCKNUT", "NUT"),
    ("LOOKNUT", "NUT"),
    ("NUT", "NUT"),
    ("DIESHOE", "SHOE"),
    ("SHOE", "SHOE"),
    ("GRIPPER", "GRIPPER"),
    ("EJECTOR", "EJECTOR"),
    ("BUSH", "บูช"),
    ("บูช", "บูช"),
    ("STRIPPER", "RING"),
    ("KNOCK", "RING"),
    ("SLEEVE", "CORE"),
    ("RING", "RING"),
    ("DIESET", "DIE SET"),
    ("PLATE", "PLATE"),
    ("COMPRESS", "PLATE"),
    ("STRIP", "PLATE"),
    ("UPPER", "PLATE"),
    ("LOWER", "PLATE"),
    ("BASE", "PLATE"),
    ("ROLLER", "PLATE"),
    ("SUPPORT", "PLATE"),
    ("BALANCE", "PLATE"),
    ("CORE", "CORE"),
    ("PUNCHCENTER", "CORE"),
    ("SHAFT", "เพลา"),
    ("เพลา", "เพลา"),
    ("GEAR", "เฟือง"),
    ("ดุม", "ดุม"),
    ("M16", "สกรู"),
    ("น็อต", "สกรู"),
    ("สลัก", "สลัก"),
    ("ก้านธง", "ทั่วไป"),
    ("ลิ่ม", "ลิ่ม"),
    ("สายพาน", "สายพาน"),
    ("ฝัก", "ฝัก"),
    ("พูลเลย์", "พูลเลย์"),
]

SHEET_GROUP = {"DIEP1": "DIE CUTTER", "DIEP3": "DIE CUTTER",
               "PUNCHP1": "PUNCH CUTTER", "PUNCHP3": "PUNCH CUTTER"}


def norm_job(job):
    return re.sub(r"\s+", "", str(job or "")).upper()


def classify_job(job_name):
    """คืน (sheet|None, group) — sheet จาก exact match ของแต่ละส่วนที่แยกด้วย + เท่านั้น"""
    parts = [norm_job(p) for p in str(job_name or "").split("+")]
    sheet = None
    for p in parts:
        if p in JOB_TO_SHEET:
            sheet = JOB_TO_SHEET[p]
            break
    key = norm_job(job_name)
    if not key:
        return sheet, "ไม่ระบุ"
    if key in JOB_GROUP_EXPLICIT:
        return sheet, JOB_GROUP_EXPLICIT[key]
    if sheet in SHEET_GROUP:
        # งานคู่ (A+B) ใช้ group ของส่วนแรกที่เจอชีท
        return sheet, SHEET_GROUP[sheet]
    for kw, g in GROUP_KEYWORDS:
        if kw in key:
            return None, g
    return None, "อื่น ๆ"


def parse_size_id(size):
    """แยก base กับ ID: '129ID92' -> ('129', 'ID92') ; '211OD115' -> ('211OD115', None)"""
    t = (size or "").replace(" ", "")
    m = re.search(r"ID\d+$", t, re.I)
    if m and t[:m.start()]:
        return t[:m.start()], m.group(0).upper()
    return t, None


def split_size_no(s):
    """S.129 ID92 No.6613 -> (129ID92, [6613]) ; S.307-2H No.6418,6607 -> (307-2H, [...])
    เก็บไว้ใช้ที่เดียว (split_groups คืนกลุ่มแรกให้เหมือนเดิม)"""
    gs = split_groups(s)
    if not gs:
        return None, []
    return gs[0]["size"], gs[0]["nos"]


def split_groups(s):
    """แยก size_ref เป็นกลุ่มๆ: [{'size':..,'nos':[..],'unnumbered':n}]
    เช่น 'S.300-2H/Fสั้น No.6502,6515 / S.300-2H/Fยาว No.6414' -> 2 กลุ่ม
    No. หลายตัวคั่นด้วย , / ; หรือช่องว่าง (เช่น 6805/6809/6814/6815)
    ตัดหมายเหตุวงเล็บที่มี No. (เช่น '(No.6503 รับแล้ว)') และนับ 'ไม่มีNo. X ตัว'"""
    if not s:
        return []
    t = str(s)
    # ตัดวงเล็บหมายเหตุที่มี No. อยู่ข้างใน
    t = re.sub(r"\([^()]*No\.[^()]*\)", " ", t)
    # นับตัวที่ไม่ระบุ No.
    unnumbered = 0
    for m in re.finditer(r"ไม่มี\s*No\.\s*(\d+)\s*ตัว", t):
        unnumbered += int(m.group(1))
    t = re.sub(r"ไม่มี\s*No\.\s*\d*\s*ตัว?", " ", t)
    groups = []
    # แต่ละกลุ่มเริ่มที่ S. จบก่อน S. ถัดไป (กัน S. ในข้อความอื่นด้วยการบังคับ No. ตาม)
    for m in re.finditer(r"S\.\s*(.*?)\s*No\.\s*((?:(?!S\.).)*?)(?=\s*S\.|$)", t):
        size = m.group(1).strip()
        nos_raw = m.group(2).strip()
        nos = [x.strip() for x in re.split(r"[,/;\s]+", nos_raw) if x.strip()]
        nos = [x for x in nos if re.fullmatch(r"[A-Za-z0-9.\-]+", x)]
        if size:
            groups.append({"size": size, "nos": nos, "unnumbered": 0})
    if unnumbered and groups:
        groups[-1]["unnumbered"] = groups[-1].get("unnumbered", 0) + unnumbered
    elif unnumbered:
        groups.append({"size": None, "nos": [], "unnumbered": unnumbered})
    if not groups:
        # ไม่มี No. เลย แต่มี S. อยู่ -> กลุ่มเปล่าสำหรับ size_only
        m = re.search(r"S\.([A-Za-z0-9.\-/ ]+)", t)
        if m and m.group(1).strip():
            groups.append({"size": m.group(1).strip(), "nos": [], "unnumbered": 0})
    return groups


def match_group(sheet, size, nos, idx, base_idx, sheet_has_id, sizes):
    """เทียบ No. กลุ่มเดียวกับทะเบียน คืน (linked, ambiguous, fuzzy, size_hit)
    กฎปลอดภัย (ID/OD-aware): ตรงเป๊ะ > ID ตรงกัน > ชีทที่ไม่แยก ID (DIE-style) >
    slip ไม่ระบุ ID แต่เหลือตัวเลือกเดียวและไม่มี ID > นอกนั้นกำกวม/ไม่พบ (ห้ามเดา)"""
    size_n = (size or "").replace(" ", "")
    linked, ambiguous, fuzzy = [], {}, 0
    if not size:
        return linked, ambiguous, fuzzy, False
    s_base, s_id = parse_size_id(size_n)
    size_hit = (sheet, size_n) in sizes or any(
        s == size_n or parse_size_id(s)[0] == s_base for (sh, s) in sizes if sh == sheet)
    if not nos:
        return linked, ambiguous, fuzzy, size_hit
    for no in nos:
        no = (no or "").strip()
        if not no:
            continue
        code = idx.get((sheet, size_n, no))
        if code:
            linked.append(code)
            continue
        cands = base_idx.get((sheet, s_base, no), [])
        done = False
        if s_id:
            same = [c for c in cands if c["id"] == s_id]
            if same:
                linked.append(same[0]["code"])
                if same[0]["size"] != size_n:
                    fuzzy += 1  # ตัวหนังสือต่างแต่ ID ตรง
                done = True
            elif sheet not in sheet_has_id and len(cands) == 1 and cands[0]["id"] is None:
                linked.append(cands[0]["code"])
                fuzzy += 1  # registry ชีทนี้ไม่แยก ID
                done = True
            # ไม่งั้น MISS: ID ไม่ตรง = คนละตัว ห้ามเดา -> ไป prefix fallback
        else:
            # slip ไม่ระบุ ID: ผูกเองได้เฉพาะ registry ก็ไม่มี ID (แน่ใจว่าเป็นตัวเดียวกัน)
            # ถ้ามีตัวเลือกติด ID -> ให้คนเลือกเอง (ambiguous) ห้ามเดา
            no_id = [c for c in cands if c["id"] is None]
            if len(cands) == 1 and len(no_id) == 1:
                linked.append(cands[0]["code"])
                done = True
            elif cands:
                ambiguous[no] = [c["code"] for c in cands]
                done = True
            # ไม่มีเลย -> ไป prefix fallback
        if not done and len(size_n) >= 3:
            # fallback สุดท้าย: SIZE อ่านไม่ครบ (เช่น '300-1H/F…')
            # เสนอตัวที่ขึ้นต้นตรงกันให้คนเลือกเอง (manual pick เท่านั้น ห้าม auto)
            pref = sorted({s for (sh, s) in sizes
                           if sh == sheet and s.startswith(size_n) and s != size_n})
            opts = []
            for s in pref:
                cd = idx.get((sheet, s, no))
                if cd and cd not in opts:
                    opts.append(cd)
            if opts:
                cur = ambiguous.setdefault(no, [])
                for cd in opts:
                    if cd not in cur and cd not in linked:
                        cur.append(cd)
    return linked, ambiguous, fuzzy, size_hit


def link_repairs(repairs, molds):
    """ผูกใบซ่อมกับตัวแม่พิมพ์ (ID/OD-aware ปลอดภัย):
    - ตรงเป๊ะ (sheet/size/no) -> ผูกทันที
    - slip มี ID, registry มี ID เดียวกัน -> ผูก (note ถ้าตัวหนังสือต่าง)
    - slip มี ID, registry ชีทนี้ไม่แยก ID เลย -> ผูกฐานเดียวกัน (DIE-style)
    - slip ไม่มี ID: ผูกได้เฉพาะกรณีเหลือตัวเลือกเดียวและ registry ก็ไม่มี ID -> ไม่งั้น ambiguous
    - นอกนั้น MISS (ห้ามเดา — ID/OD ต่างกันคือคนละตัว)
    คืน link {sheet,size,no_list,linked_codes,link_status,link_note,group,ambiguous}"""
    idx = {}
    base_idx = {}
    sheet_has_id = set()
    for m in molds:
        sz = (m["size"] or "").replace(" ", "")
        idx[(m["sheet"], sz, (m["no"] or "").strip())] = m["code"]
        base, mid = parse_size_id(sz)
        base_idx.setdefault((m["sheet"], base, (m["no"] or "").strip()), []).append(
            {"size": sz, "id": mid, "code": m["code"]})
        if mid:
            sheet_has_id.add(m["sheet"])
    sizes = {(m["sheet"], (m["size"] or "").replace(" ", "")) for m in molds}
    orphans = []
    for r in repairs:
        parts = [p for p in str(r.get("job_name") or "").split("+")]
        sheet = None
        for p in parts:
            if norm_job(p) in JOB_TO_SHEET:
                sheet = JOB_TO_SHEET[norm_job(p)]
                break
        group = classify_job(r.get("job_name"))[1]
        groups = split_groups(r.get("size_ref"))
        linked = []
        ambiguous = {}
        fuzzy = 0
        size_hit_any = False
        grp_out = []
        note = None
        if sheet is None:
            status, note = "orphan", "งานนอกทะเบียนสต็อก"
        elif not groups:
            status, note = "orphan", "อ่าน SIZE ไม่ได้"
        else:
            all_nos = []
            for g in groups:
                g_linked, g_amb, g_fuzzy, g_hit = match_group(
                    sheet, g["size"], g["nos"], idx, base_idx, sheet_has_id, sizes)
                if g_hit:
                    size_hit_any = True
                linked.extend(g_linked)
                fuzzy += g_fuzzy
                for k, v in g_amb.items():
                    ambiguous.setdefault(k, [])
                    for c in v:
                        if c not in ambiguous[k]:
                            ambiguous[k].append(c)
                grp_out.append({"size": g["size"], "nos": g["nos"],
                                "linked": g_linked, "unnumbered": g["unnumbered"]})
                all_nos.extend([n for n in g["nos"]])
            nos = all_nos
            size = groups[0]["size"] if groups else None
            unnum = sum(g.get("unnumbered", 0) for g in grp_out)
            if len(linked) == len(nos) and nos:
                status = "linked"
                if fuzzy:
                    note = f"เทียบ SIZE แบบยืดหยุ่น ({fuzzy} ตัว)"
            elif not nos and size_hit_any:
                status = "size_only"  # ระบุถึงระดับ SIZE แต่ไม่ลง No.
            elif linked or ambiguous:
                got = len(linked)
                tot = len(nos)
                amb = sum(1 for no in nos if no.strip() in ambiguous)
                status, note = "partial", f"พบ {got}/{tot} ตัว" + (f" · กำกวม {amb} No. (เลือกเองตอนรับกลับ)" if amb else "")
            else:
                shown = groups[0]["size"] if groups else size
                status, note = "orphan", f"ไม่พบ {sheet}/{shown} ในสต็อก"
            if unnum and status in ("linked", "partial"):
                note = ((note + " · ") if note else "") + f"มี {unnum} ตัวไม่ระบุ No."
        r["link"] = {"sheet": sheet, "size": groups[0]["size"] if groups else None,
                     "no_list": [n for g in grp_out for n in g["nos"]],
                     "group": group, "linked_codes": linked, "link_status": status,
                     "link_note": note, "ambiguous": ambiguous, "groups": grp_out}
        if status != "linked":
            orphans.append({"source_no": r["source_no"], "job_name": r.get("job_name"), "group": group,
                            "size_ref": r.get("size_ref"), "link_status": status, "link_note": note})
    return orphans


def carry_registered(molds):
    """ดึงของที่ลงทะเบียนผ่านระบบ (wf.mold_static) กลับมาต่อท้าย — กัน import รอบใหม่ทับหาย"""
    wf_db = os.path.join(ROOT, "data", "wf.db")
    if not os.path.exists(wf_db):
        return 0
    try:
        con = sqlite3.connect(wf_db)
        tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if "mold_static" not in tables:
            con.close()
            return 0
        rows = con.execute("SELECT code, data FROM mold_static").fetchall()
        live = {r[0]: (r[1], r[2]) for r in
                con.execute("SELECT code, status, machine FROM mold_live")}
        con.close()
    except Exception as e:  # noqa: BLE001
        print(f"  ! อ่านของลงทะเบียนเดิมไม่ได้: {e}")
        return 0
    have = {m["code"] for m in molds}
    max_block = {}
    for m in molds:
        max_block[m["sheet"]] = max(max_block.get(m["sheet"], 0), m.get("block", 0))
    added = 0
    for code, data in rows:
        if code in have:
            continue
        try:
            e = json.loads(data)
        except ValueError:
            continue
        st, mc = live.get(code, ("spare", None))
        e["status"] = st if st in STATUS_TH else "spare"
        e["status_th"] = STATUS_TH[e["status"]]
        e["machine"] = mc
        used = {(m["sheet"], m.get("block")) for m in molds}
        if (e.get("sheet"), e.get("block")) in used:
            max_block[e["sheet"]] = max_block.get(e["sheet"], 0) + 1
            e["block"], e["seq"] = max_block[e["sheet"]], 1
        molds.append(e)
        added += 1
    if added:
        print(f"  ต่อของลงทะเบียนเดิม {added} ตัว (ไม่ทับ)")
    return added


# ชื่อกลุ่มใน sheet รหัสแบบ -> ชีทสต็อก (None = กลุ่มนอกทะเบียน แต่เก็บไว้เป็นข้อมูลจริง)
PATTERN_GROUP_SHEET = {
    "PRESSURE RING P2": None,
    "REDRAW PRESSURE P2": None,
    "DIE CUTTER P1": "DIEP1",
    "PUNCH CUTTER P1": "PUNCHP1",
    "PUNCH CUTTER P3": "PUNCHP3",
    "DIE CUTTER P3": "DIEP3",
}


def norm_pattern_size(label):
    """ตัดวงเล็บคำอธิบาย + ช่องว่าง: '109 (209.5/208x100)' -> '109', '300 OD115/Fกลาง' -> '300OD115/Fกลาง'"""
    t = str(label or "")
    t = t.split("(")[0].replace(" ", "").strip()
    return t


def parse_patterns(wb):
    """อ่าน master รหัสแบบจาก sheet 'รหัสแบบ' -> [{group, sheet, size_label, size_key, pattern}]
    โครงชีท: คู่คอลัมน์ (SIZE | รหัสแบบ), กลุ่ม PUNCH P1 มีคอลัมน์ ID คั่นกลาง (SIZE | ID | รหัสแบบ)"""
    if "รหัสแบบ" not in wb.sheetnames:
        print("  ! ไม่มีชีท รหัสแบบ")
        return []
    ws = wb["รหัสแบบ"]
    header = [(str(ws.cell(row=1, column=c).value or "").strip()) for c in range(1, ws.max_column + 1)]
    groups = []  # (title, size_col, id_col|None, pat_col) 1-based
    for j, h in enumerate(header):
        if h != "รหัสแบบ":
            continue
        pat = j + 1
        idc = size = title = None
        if header[j - 1] == "ID":
            idc = j  # 1-based col of ID
            size = j - 1
            k = j - 2
        else:
            size = j
            k = j - 1
        while k >= 0 and header[k] in ("", "รหัสแบบ", "ID"):
            k -= 1
        title = header[k] if k >= 0 else ""
        groups.append((title, size, idc, pat))
    out = []
    for r in range(2, ws.max_row + 1):
        for title, size_c, id_c, pat_c in groups:
            label = ws.cell(row=r, column=size_c).value
            pat = ws.cell(row=r, column=pat_c).value
            if label is None or str(label).strip() == "":
                continue
            if pat is None or str(pat).strip() in ("", "-"):
                continue
            base = norm_pattern_size(label)
            if not base:
                continue
            idv = str(ws.cell(row=r, column=id_c).value or "").strip() if id_c else ""
            key = f"{base}ID{idv}" if idv else base
            tkey = re.sub(r"\s+", " ", title).strip().upper()
            sheet = None
            for name, s in PATTERN_GROUP_SHEET.items():
                if tkey == name.upper():
                    sheet = s
                    break
            out.append({"group": title, "sheet": sheet, "size_label": str(label).strip(),
                        "size_key": key, "pattern": str(pat).strip()})
    # ตัดซ้ำ
    seen, uniq = set(), []
    for e in out:
        k = (e["group"], e["size_key"], e["pattern"])
        if k not in seen:
            seen.add(k)
            uniq.append(e)
    print(f"  รหัสแบบ: {len(uniq)} รายการ จาก {len(groups)} กลุ่ม")
    return uniq


def link_patterns(patterns, molds):
    """รายงานว่า size_key ตรงกับ SIZE สต็อกแค่ไหน (exact ก่อน, prefix รอง)"""
    stock = {}
    for m in molds:
        stock.setdefault(m["sheet"], set()).add(m["size"])
    exact = prefix = miss = 0
    for p in patterns:
        sizes = stock.get(p["sheet"], set()) if p["sheet"] else set()
        if p["size_key"] in sizes:
            exact += 1
        elif p["sheet"] and any(s.startswith(p["size_key"]) for s in sizes):
            prefix += 1
        else:
            miss += 1
    print(f"  ผูก SIZE สต็อก: ตรงเป๊ะ {exact} / นำหน้า {prefix} / ไม่เจอ {miss}")
    return {"exact": exact, "prefix": prefix, "miss": miss}


def load_wf_state():
    """อ่านสถานะระบบงานจริง: ของที่แตะแล้ว / tombstone / meta / ค่า live+DB
    นโยบาย: รายการที่ระบบแตะแล้ว -> ฐานข้อมูลระบบชนะ Excel; ที่เหลือ Excel ชนะ"""
    st = {"touched_molds": set(), "touched_papers": set(), "tombstones": set(),
          "meta": {}, "db_molds": {}, "db_papers": {}, "dims": {}}
    wf_db = os.path.join(ROOT, "data", "wf.db")
    try:
        con = sqlite3.connect(f"file:{wf_db}?mode=ro", uri=True)
        con.row_factory = sqlite3.Row
        tables = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if "movements" in tables:
            for (c,) in con.execute("SELECT DISTINCT mold_code FROM movements"):
                st["touched_molds"].add(c)
            for (d,) in con.execute("SELECT DISTINCT doc_no FROM movements"):
                d = str(d or "")
                st["touched_papers"].add(d[:-5] if d.endswith("-VOID") else d)
        if "mold_static" in tables:
            for (c,) in con.execute("SELECT code FROM mold_static"):
                st["touched_molds"].add(c)
        if "docs" in tables:
            for (d,) in con.execute("SELECT doc_no FROM docs WHERE doc_type='P4R'"):
                st["touched_papers"].add(d)
        if "repair_meta" in tables:
            for r in con.execute("SELECT * FROM repair_meta"):
                r = dict(r)
                sno = r.pop("source_no")
                if r.pop("deleted", 0):
                    st["tombstones"].add(sno)
                else:
                    st["meta"][sno] = {k: v for k, v in r.items() if v is not None}
                    st["touched_papers"].add(sno)  # ใบที่เคยแก้ในระบบ: DB+meta ชนะ Excel
        if "inspections" in tables:
            # dims ล่าสุดรายตัว (จากใบตรวจรับ) ชนะค่า Excel
            for r in con.execute("SELECT mold_code,measured FROM inspections WHERE id IN"
                                 " (SELECT MAX(id) FROM inspections GROUP BY mold_code)"):
                try:
                    d = json.loads(r[1] or "{}")
                    if isinstance(d, dict) and d:
                        st["dims"][r[0]] = d
                        st["touched_molds"].add(r[0])
                except ValueError:
                    continue
        con.close()
    except Exception as e:  # noqa: BLE001
        print(f"  ! อ่าน wf.db ไม่ได้: {e}")
    try:
        app = sqlite3.connect(f"file:{DB_PATH}?mode=ro", uri=True)
        app.row_factory = sqlite3.Row
        for r in app.execute("SELECT code,status,machine FROM molds"):
            st["db_molds"][r["code"]] = (r["status"], r["machine"])
        for r in app.execute("SELECT source_no,job_name,status,ordered_date FROM repair_orders"):
            st["db_papers"][r["source_no"]] = dict(r)
        app.close()
    except Exception as e:  # noqa: BLE001
        print(f"  ! อ่าน app.db ไม่ได้: {e}")
    print(f"  สถานะระบบ: แตะแม่พิมพ์ {len(st['touched_molds'])} ตัว / ใบที่แตะ {len(st['touched_papers'])} ใบ / tombstone {len(st['tombstones'])} ใบ")
    return st


META_JSON_MAP = {"detail": "detail", "pattern_code": "pattern_code", "requester": "requester",
                 "machine": "machine", "department": "department", "remark": "remark",
                 "qty": "qty", "unit": "unit", "priority": "priority", "due_date": "due_date",
                 "job_name": "job_name", "ordered_date": "ordered_date"}


def apply_system_overlay(molds, repairs, orphans, st):
    """ระบบชนะ Excel สำหรับรายการที่แตะแล้ว; ตัด tombstone ทิ้ง"""
    for m in molds:
        if m["code"] in st["touched_molds"] and m["code"] in st["db_molds"]:
            s, mc = st["db_molds"][m["code"]]
            m["status"] = s
            m["status_th"] = STATUS_TH.get(s, m.get("status_th"))
            m["machine"] = mc
        if m["code"] in st["dims"]:
            m["dims"] = st["dims"][m["code"]]
    out = []
    for r in repairs:
        sno = r["source_no"]
        if sno in st["tombstones"]:
            continue
        if sno in st["touched_papers"]:
            if sno in st["db_papers"]:
                d = st["db_papers"][sno]
                r["status"] = d["status"]
                r["job_name"] = d["job_name"]
                r["ordered_date"] = d["ordered_date"]
            for jk, jv in META_JSON_MAP.items():
                if sno in st["meta"] and jk in st["meta"][sno]:
                    r[jv] = st["meta"][sno][jk]
        out.append(r)
    orphans = [o for o in orphans if o["source_no"] not in st["tombstones"]]
    return out, orphans


def reconcile_open_repairs(molds, repairs, st):
    """สถานะต้องสอดคล้องกับใบซ่อม: แม่พิมพ์ที่ผูกกับใบซ่อม open ต้องเป็น repair.

    เรียกหลัง apply_system_overlay / ก่อน write_json+write_db เพื่อให้ JSON และ DB
    ตรงกันทุกครั้งที่ import ใหม่ (กันไม่ให้ชีทสต็อกที่ลืมอัปเดตทำให้ข้อมูลเพี้ยนอีก)
    - ข้ามแม่พิมพ์ที่แตะในระบบแล้ว (touched_molds: ระบบชนะ) แต่พิมพ์เตือนให้ตรวจ
    """
    open_codes = {}
    for r in repairs:
        if r.get("status") == "open":
            for c in (r.get("link") or {}).get("linked_codes", []):
                open_codes.setdefault(c, []).append(r["source_no"])
    by_code = {m["code"]: m for m in molds}
    fixed, skipped = 0, []
    for code, papers in open_codes.items():
        m = by_code.get(code)
        if not m or m["status"] == "repair":
            continue
        if code in st["touched_molds"]:
            skipped.append(code)
            continue
        m["status"] = "repair"
        m["status_th"] = STATUS_TH.get("repair", "ส่งซ่อม")
        fixed += 1
    print(f"  reconcile: แก้สถานะเป็น repair {fixed} ตัว (ผูกกับใบซ่อมค้าง)")
    if skipped:
        print(f"  ! ข้าม {len(skipped)} ตัวที่แตะในระบบแล้ว ควรตรวจเอง: {', '.join(skipped[:10])}")


def write_json(molds, repairs):
    os.makedirs(DATA_DIR, exist_ok=True)
    with io.open(os.path.join(DATA_DIR, "molds.json"), "w", encoding="utf-8") as fh:
        json.dump(molds, fh, ensure_ascii=False)
    with io.open(os.path.join(DATA_DIR, "repairs.json"), "w", encoding="utf-8") as fh:
        json.dump(repairs, fh, ensure_ascii=False)
    by_type = {}
    for m in molds:
        t = by_type.setdefault(m["sheet"], {"total": 0, "spare": 0, "installed": 0, "repair": 0, "retired": 0})
        t["total"] += 1
        t[m["status"]] += 1
    meta = {
        "imported_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "molds_total": len(molds),
        "by_type": by_type,
        "repairs_total": len(repairs),
        "repairs_open": sum(1 for r in repairs if r["status"] == "open"),
        "repairs_linked": sum(1 for r in repairs if r.get("link", {}).get("link_status") == "linked"),
        "repairs_orphan": sum(1 for r in repairs if r.get("link", {}).get("link_status") != "linked"),
    }
    with io.open(os.path.join(DATA_DIR, "meta.json"), "w", encoding="utf-8") as fh:
        json.dump(meta, fh, ensure_ascii=False, indent=1)
    print(f"  เขียน JSON: molds={len(molds)} repairs={len(repairs)} -> {DATA_DIR}")
    return meta


def write_db(molds, repairs, st):
    if not os.path.exists(DB_PATH):
        print(f"  ! ไม่พบ {DB_PATH} ข้ามการเขียน DB")
        return
    bak = DB_PATH + ".bak"
    shutil.copy2(DB_PATH, bak)
    print(f"  สำรอง DB -> {bak}")
    con = sqlite3.connect(DB_PATH, timeout=30.0)
    now = datetime.now().astimezone().isoformat(timespec="seconds")
    cur = con.cursor()
    cur.execute("DELETE FROM molds WHERE code LIKE 'DEMO-%'")
    cur.execute("DELETE FROM repair_orders WHERE source_no LIKE 'P4R-10-000%' OR source_no LIKE 'P4R-09-0258'")
    # tombstone: ใบที่ถูกลบในระบบต้องไม่กลับมา
    if st["tombstones"]:
        cur.execute(f"DELETE FROM repair_orders WHERE source_no IN ({','.join('?' * len(st['tombstones']))})",
                    tuple(st["tombstones"]))
        print(f"  ตัด tombstone {len(st['tombstones'])} ใบ")
    new_m, upd_m = 0, 0
    for m in molds:
        if m["code"] in st["touched_molds"]:
            continue  # ระบบชนะ: คงค่า DB ไว้ (JSON ถูก overlay แล้ว)
        cur.execute("INSERT OR IGNORE INTO molds(code, tooling_type, size, status, machine, created_at)"
                    " VALUES(?,?,?,?,?,?)",
                    (m["code"], m["tooling_type"], m["size"], m["status"], m["machine"], now))
        if cur.rowcount:
            new_m += 1
        else:  # ไม่แตะในระบบ: Excel ชนะ อัปเดตค่าตามไฟล์
            cur.execute("UPDATE molds SET tooling_type=?, size=?, status=?, machine=? WHERE code=?",
                        (m["tooling_type"], m["size"], m["status"], m["machine"], m["code"]))
            upd_m += cur.rowcount
    new_r, upd_r = 0, 0
    for r in repairs:
        if r["source_no"] in st["touched_papers"]:
            continue  # ระบบชนะ: คงค่า DB ไว้ (JSON ถูก overlay แล้ว)
        cur.execute("INSERT OR IGNORE INTO repair_orders(source_no, job_name, status, ordered_date, created_at)"
                    " VALUES(?,?,?,?,?)",
                    (r["source_no"], r["job_name"] or r["source_no"], r["status"], r["ordered_date"], now))
        if cur.rowcount:
            new_r += 1
        else:
            cur.execute("UPDATE repair_orders SET job_name=?, status=?, ordered_date=? WHERE source_no=?",
                        (r["job_name"] or r["source_no"], r["status"], r["ordered_date"], r["source_no"]))
            upd_r += cur.rowcount
    con.commit()
    n_m = cur.execute("SELECT COUNT(*) FROM molds").fetchone()[0]
    n_r = cur.execute("SELECT COUNT(*) FROM repair_orders").fetchone()[0]
    con.close()
    print(f"  DB ปัจจุบัน: molds={n_m} (ใหม่ {new_m}/อัปเดต {upd_m}) repair_orders={n_r} (ใหม่ {new_r}/อัปเดต {upd_r})")


def main():
    xls = find_excel()
    print(f"ไฟล์: {os.path.basename(xls)}")
    wb = openpyxl.load_workbook(xls, data_only=True, read_only=False)
    print("อ่านสต็อก...")
    molds = parse_stock(wb)
    print("อ่านใบซ่อม...")
    repairs = parse_repairs(wb)
    print("ผูกใบซ่อมกับสต็อก...")
    orphans = link_repairs(repairs, molds)
    print(f"  ผูกได้ {sum(1 for r in repairs if r['link']['link_status']=='linked')}/{len(repairs)} ใบ "
          f"(บางส่วน {sum(1 for r in repairs if r['link']['link_status']=='partial')}, "
          f"ผูกไม่ได้ {sum(1 for r in repairs if r['link']['link_status']=='orphan')})")
    print("เขียนไฟล์...")
    carry_registered(molds)
    st = load_wf_state()
    repairs, orphans = apply_system_overlay(molds, repairs, orphans, st)
    print("ตรวจสอบสถานะกับใบซ่อมค้าง...")
    reconcile_open_repairs(molds, repairs, st)
    print("อ่าน master รหัสแบบ...")
    patterns = parse_patterns(wb)
    with io.open(os.path.join(DATA_DIR, "patterns.json"), "w", encoding="utf-8") as fh:
        json.dump(patterns, fh, ensure_ascii=False)
    pat_link = link_patterns(patterns, molds)
    with io.open(os.path.join(DATA_DIR, "orphans.json"), "w", encoding="utf-8") as fh:
        json.dump(orphans, fh, ensure_ascii=False)
    meta = write_json(molds, repairs)
    meta["patterns_total"] = len(patterns)
    meta["patterns_link"] = pat_link
    with io.open(os.path.join(DATA_DIR, "meta.json"), "w", encoding="utf-8") as fh:
        json.dump(meta, fh, ensure_ascii=False, indent=1)
    write_db(molds, repairs, st)
    print(f"เสร็จ: แม่พิมพ์ {meta['molds_total']} ตัว, ใบซ่อม {meta['repairs_total']} ใบ, รหัสแบบ {len(patterns)} รายการ")


if __name__ == "__main__":
    sys.exit(main())
